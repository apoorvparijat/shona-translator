#!/usr/bin/env python3
"""
Translation Verifier - Analyzes input and output DOCX files to generate translation quality reports
"""

import os
import csv
import re
import time
import logging
from typing import List, Dict, Tuple, Optional
from docx import Document
from dotenv import load_dotenv
from difflib import SequenceMatcher

# Try to import OpenAI, but make it optional
try:
    import openai
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False
    logger = logging.getLogger(__name__)
    logger.warning("OpenAI library not available. Using basic confidence scoring only.")

# Load environment variables
load_dotenv()

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class TranslationVerifier:
    """Verifies translation quality by comparing input and output DOCX files"""
    
    def __init__(self):
        self.openai_client = None
        
        # Initialize OpenAI if API key is available and library is installed
        if OPENAI_AVAILABLE and os.getenv('OPENAI_API_KEY'):
            try:
                openai.api_key = os.getenv('OPENAI_API_KEY')
                self.openai_client = openai.OpenAI()
                logger.info("OpenAI client initialized for verification")
            except Exception as e:
                logger.warning(f"Failed to initialize OpenAI client: {e}")
                self.openai_client = None
        else:
            if not OPENAI_AVAILABLE:
                logger.warning("OpenAI library not available. Using basic confidence scoring only.")
            else:
                logger.warning("OpenAI API key not found. Using basic confidence scoring only.")
    
    def extract_text_from_docx(self, file_path: str) -> List[str]:
        """Extract all text lines from a DOCX file, including tables"""
        text_lines = []
        
        try:
            doc = Document(file_path)
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:  # Only add non-empty lines
                    text_lines.append(text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            text = paragraph.text.strip()
                            if text:  # Only add non-empty lines
                                text_lines.append(text)
            
            logger.info(f"Extracted {len(text_lines)} text lines from {file_path}")
            return text_lines
            
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {e}")
            return []
    
    def calculate_basic_confidence(self, original: str, translated: str) -> int:
        """Calculate basic confidence score based on heuristics (0-100)"""
        if not original.strip() or not translated.strip():
            return 0
        
        # If texts are identical, likely untranslated
        if original.lower().strip() == translated.lower().strip():
            return 10  # Very low confidence - probably not translated
        
        # Basic length ratio check
        length_ratio = len(translated) / len(original) if len(original) > 0 else 0
        length_score = 100 if 0.5 <= length_ratio <= 2.0 else max(0, 100 - abs(length_ratio - 1) * 50)
        
        # Character diversity check (Shona uses different character patterns than English)
        english_chars = set('abcdefghijklmnopqrstuvwxyz')
        original_chars = set(original.lower())
        translated_chars = set(translated.lower())
        
        english_ratio_original = len(original_chars & english_chars) / len(original_chars) if original_chars else 0
        english_ratio_translated = len(translated_chars & english_chars) / len(translated_chars) if translated_chars else 0
        
        # Good translation should have some character pattern change
        diversity_score = 100 if english_ratio_translated < english_ratio_original else 70
        
        # Word count similarity
        original_words = len(original.split())
        translated_words = len(translated.split())
        word_ratio = translated_words / original_words if original_words > 0 else 0
        word_score = 100 if 0.7 <= word_ratio <= 1.5 else max(0, 100 - abs(word_ratio - 1) * 30)
        
        # Combine scores with weights
        basic_confidence = int(
            length_score * 0.3 + 
            diversity_score * 0.4 + 
            word_score * 0.3
        )
        
        return max(0, min(100, basic_confidence))
    
    def calculate_ai_confidence(self, original: str, translated: str) -> Optional[int]:
        """Calculate AI-based confidence score using OpenAI (0-100)"""
        if not self.openai_client:
            return None
        
        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4",
                messages=[
                    {
                        "role": "system",
                        "content": """You are a translation quality assessor specializing in English to Shona translations.
                        Shona is a Bantu language spoken primarily in Zimbabwe.
                        
                        Rate the translation quality on a scale of 0-100 where:
                        - 90-100: Excellent translation (accurate, natural, culturally appropriate)
                        - 75-89: Good translation (mostly accurate with minor issues)
                        - 60-74: Acceptable translation (understandable but some errors)
                        - 40-59: Poor translation (significant errors, unclear meaning)
                        - 20-39: Very poor translation (major errors, meaning lost)
                        - 0-19: Failed translation (untranslated, gibberish, or completely wrong)
                        
                        Consider: accuracy, naturalness, cultural appropriateness, grammar.
                        Respond with ONLY the numeric score (0-100)."""
                    },
                    {
                        "role": "user",
                        "content": f"Original English: {original}\nShona Translation: {translated}\n\nQuality score (0-100):"
                    }
                ],
                temperature=0.1,
                max_tokens=10
            )
            
            # Extract numeric score from response
            score_text = response.choices[0].message.content.strip()
            score = int(re.search(r'\d+', score_text).group()) if re.search(r'\d+', score_text) else 0
            return max(0, min(100, score))
            
        except Exception as e:
            logger.error(f"Error calculating AI confidence: {e}")
            return None
    
    def align_texts(self, original_lines: List[str], translated_lines: List[str]) -> List[Tuple[str, str]]:
        """Align original and translated text lines using sequence matching"""
        aligned_pairs = []
        
        # If lengths are equal, assume 1:1 mapping
        if len(original_lines) == len(translated_lines):
            for orig, trans in zip(original_lines, translated_lines):
                aligned_pairs.append((orig, trans))
        else:
            # Use sequence matching for different lengths
            matcher = SequenceMatcher(None, original_lines, translated_lines)
            
            for tag, i1, i2, j1, j2 in matcher.get_opcodes():
                if tag == 'equal' or tag == 'replace':
                    # Pair up lines in this segment
                    orig_segment = original_lines[i1:i2]
                    trans_segment = translated_lines[j1:j2]
                    
                    max_len = max(len(orig_segment), len(trans_segment))
                    for k in range(max_len):
                        orig = orig_segment[k] if k < len(orig_segment) else ""
                        trans = trans_segment[k] if k < len(trans_segment) else ""
                        aligned_pairs.append((orig, trans))
                        
                elif tag == 'delete':
                    # Original lines with no translation
                    for k in range(i1, i2):
                        aligned_pairs.append((original_lines[k], ""))
                        
                elif tag == 'insert':
                    # Translated lines with no original
                    for k in range(j1, j2):
                        aligned_pairs.append("", translated_lines[k])
        
        return aligned_pairs
    
    def generate_verification_report(self, input_file: str, output_file: str, report_file: str) -> bool:
        """Generate CSV verification report comparing input and output files"""
        try:
            logger.info(f"Starting verification of {input_file} -> {output_file}")
            
            # Extract text from both files
            original_lines = self.extract_text_from_docx(input_file)
            translated_lines = self.extract_text_from_docx(output_file)
            
            if not original_lines:
                logger.error(f"No text extracted from input file: {input_file}")
                return False
            
            if not translated_lines:
                logger.error(f"No text extracted from output file: {output_file}")
                return False
            
            # Align the texts
            aligned_pairs = self.align_texts(original_lines, translated_lines)
            
            # Generate verification data
            verification_data = []
            for i, (original, translated) in enumerate(aligned_pairs):
                if not original and not translated:
                    continue
                
                # Calculate confidence scores
                basic_confidence = self.calculate_basic_confidence(original, translated)
                ai_confidence = self.calculate_ai_confidence(original, translated) if self.openai_client else None
                
                # Use AI confidence if available, otherwise basic confidence
                final_confidence = ai_confidence if ai_confidence is not None else basic_confidence
                
                verification_data.append({
                    'Line_Number': i + 1,
                    'Original_Text': original,
                    'Shona_Text': translated,
                    'Confidence': final_confidence,
                    'Basic_Confidence': basic_confidence,
                    'AI_Confidence': ai_confidence if ai_confidence is not None else 'N/A',
                    'Notes': self._generate_notes(original, translated, final_confidence)
                })
                
                # Add delay for AI calls to avoid rate limiting
                if self.openai_client and ai_confidence is not None:
                    time.sleep(0.5)
                
                # Progress logging
                if (i + 1) % 10 == 0:
                    logger.info(f"Processed {i + 1}/{len(aligned_pairs)} text pairs")
            
            # Write CSV report
            self._write_csv_report(verification_data, report_file)
            
            # Generate summary statistics
            self._generate_summary_stats(verification_data, report_file)
            
            logger.info(f"Verification report generated: {report_file}")
            return True
            
        except Exception as e:
            logger.error(f"Error generating verification report: {e}")
            return False
    
    def _generate_notes(self, original: str, translated: str, confidence: int) -> str:
        """Generate notes about the translation quality"""
        notes = []
        
        if confidence >= 90:
            notes.append("Excellent translation")
        elif confidence >= 75:
            notes.append("Good translation")
        elif confidence >= 60:
            notes.append("Acceptable translation")
        elif confidence >= 40:
            notes.append("Poor translation")
        elif confidence >= 20:
            notes.append("Very poor translation")
        else:
            notes.append("Failed translation")
        
        if original.lower().strip() == translated.lower().strip():
            notes.append("Possibly untranslated")
        
        if not translated.strip():
            notes.append("Missing translation")
        
        if len(translated) > len(original) * 3:
            notes.append("Translation too long")
        elif len(translated) < len(original) * 0.3:
            notes.append("Translation too short")
        
        return "; ".join(notes)
    
    def _write_csv_report(self, verification_data: List[Dict], report_file: str) -> None:
        """Write verification data to CSV file"""
        fieldnames = ['Line_Number', 'Original_Text', 'Shona_Text', 'Confidence', 'Basic_Confidence', 'AI_Confidence', 'Notes']
        
        with open(report_file, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(verification_data)
    
    def _generate_summary_stats(self, verification_data: List[Dict], report_file: str) -> None:
        """Generate and log summary statistics"""
        if not verification_data:
            return
        
        confidences = [row['Confidence'] for row in verification_data if isinstance(row['Confidence'], int)]
        
        if confidences:
            avg_confidence = sum(confidences) / len(confidences)
            min_confidence = min(confidences)
            max_confidence = max(confidences)
            
            # Quality distribution
            excellent = sum(1 for c in confidences if c >= 90)
            good = sum(1 for c in confidences if 75 <= c < 90)
            acceptable = sum(1 for c in confidences if 60 <= c < 75)
            poor = sum(1 for c in confidences if 40 <= c < 60)
            very_poor = sum(1 for c in confidences if 20 <= c < 40)
            failed = sum(1 for c in confidences if c < 20)
            
            # Write summary to separate file
            summary_file = report_file.replace('.csv', '_summary.txt')
            with open(summary_file, 'w', encoding='utf-8') as f:
                f.write("TRANSLATION VERIFICATION SUMMARY\n")
                f.write("=" * 50 + "\n\n")
                f.write(f"Total text pairs analyzed: {len(verification_data)}\n")
                f.write(f"Average confidence: {avg_confidence:.1f}%\n")
                f.write(f"Minimum confidence: {min_confidence}%\n")
                f.write(f"Maximum confidence: {max_confidence}%\n\n")
                f.write("Quality Distribution:\n")
                f.write(f"  Excellent (90-100%): {excellent} ({excellent/len(confidences)*100:.1f}%)\n")
                f.write(f"  Good (75-89%): {good} ({good/len(confidences)*100:.1f}%)\n")
                f.write(f"  Acceptable (60-74%): {acceptable} ({acceptable/len(confidences)*100:.1f}%)\n")
                f.write(f"  Poor (40-59%): {poor} ({poor/len(confidences)*100:.1f}%)\n")
                f.write(f"  Very Poor (20-39%): {very_poor} ({very_poor/len(confidences)*100:.1f}%)\n")
                f.write(f"  Failed (0-19%): {failed} ({failed/len(confidences)*100:.1f}%)\n")
            
            logger.info(f"Average translation confidence: {avg_confidence:.1f}%")
            logger.info(f"Summary statistics saved to: {summary_file}")

def main():
    """Main function to run translation verification"""
    verifier = TranslationVerifier()
    
    # Default file paths
    input_file = "collection-tools.docx"
    output_file = "collection-tools_shona.docx"
    report_file = "translation_verification_report.csv"
    
    # Check if files exist
    if not os.path.exists(input_file):
        print(f"Error: Input file '{input_file}' not found.")
        return
    
    if not os.path.exists(output_file):
        print(f"Error: Output file '{output_file}' not found.")
        print("Please run the translator first to generate the output file.")
        return
    
    print(f"🔍 Verifying translation quality...")
    print(f"📄 Input: {input_file}")
    print(f"📄 Output: {output_file}")
    print(f"📊 Report: {report_file}")
    print("\nThis may take a while for large documents...")
    
    success = verifier.generate_verification_report(input_file, output_file, report_file)
    
    if success:
        print(f"\n✅ Verification completed successfully!")
        print(f"📄 Report saved as: {report_file}")
        print(f"📊 Summary saved as: {report_file.replace('.csv', '_summary.txt')}")
        
        if verifier.openai_client:
            print("\n🤖 Used AI-powered confidence scoring")
        else:
            print("\n📊 Used basic heuristic confidence scoring")
            print("💡 Add OPENAI_API_KEY to .env for AI-powered scoring")
    else:
        print("❌ Verification failed. Check the logs for details.")

if __name__ == "__main__":
    main()